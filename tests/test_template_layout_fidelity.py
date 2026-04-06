"""Layout preservation for template-driven DOCX rendering."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Cm

pytest.importorskip("docx")

_LEN_TOL_CM = 0.03


def _write_non_a4_template(path: Path) -> None:
    """Letter-like page size with distinct margins (not SOZO A4 defaults)."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.59)
    sec.page_height = Cm(27.94)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.18)
    sec.top_margin = Cm(2.21)
    sec.bottom_margin = Cm(1.99)
    sec.header_distance = Cm(1.41)
    sec.footer_distance = Cm(1.52)
    doc.add_heading("Clinical Overview", level=1)
    doc.add_paragraph("Body under heading one.")
    doc.save(str(path))


def _assert_section_geometry_match(a, b) -> None:
    assert a.page_width.cm == pytest.approx(b.page_width.cm, abs=_LEN_TOL_CM)
    assert a.page_height.cm == pytest.approx(b.page_height.cm, abs=_LEN_TOL_CM)
    assert a.left_margin.cm == pytest.approx(b.left_margin.cm, abs=_LEN_TOL_CM)
    assert a.right_margin.cm == pytest.approx(b.right_margin.cm, abs=_LEN_TOL_CM)
    assert a.top_margin.cm == pytest.approx(b.top_margin.cm, abs=_LEN_TOL_CM)
    assert a.bottom_margin.cm == pytest.approx(b.bottom_margin.cm, abs=_LEN_TOL_CM)
    assert a.header_distance.cm == pytest.approx(b.header_distance.cm, abs=_LEN_TOL_CM)
    assert a.footer_distance.cm == pytest.approx(b.footer_distance.cm, abs=_LEN_TOL_CM)
    assert a.orientation == b.orientation


def test_render_with_layout_template_preserves_page_geometry(
    parkinsons_condition, tmp_path: Path
):
    """Output section properties match the layout template, not default A4."""
    from sozo_generator.docx.renderer import DocumentRenderer
    from sozo_generator.docx.layout import configure_page_layout
    from sozo_generator.template.doc_structure import build_document_spec
    from sozo_generator.core.enums import DocumentType, Tier

    layout_src = tmp_path / "layout_template.docx"
    _write_non_a4_template(layout_src)

    source_doc = Document(str(layout_src))
    src_sec = source_doc.sections[0]

    spec = build_document_spec(
        condition=parkinsons_condition,
        doc_type=DocumentType.EVIDENCE_BASED_PROTOCOL,
        tier=Tier.FELLOW,
    )
    out_path = tmp_path / "out.docx"
    renderer = DocumentRenderer(output_dir=str(tmp_path))
    renderer.render(spec, str(out_path), layout_template_path=layout_src)

    out_doc = Document(str(out_path))
    _assert_section_geometry_match(out_doc.sections[0], src_sec)

    default_doc = Document()
    configure_page_layout(default_doc)
    assert out_doc.sections[0].page_height.cm != pytest.approx(
        default_doc.sections[0].page_height.cm, abs=0.5
    )


def test_default_render_without_layout_template_is_a4(parkinsons_condition, tmp_path: Path):
    """Regression: no layout_template_path keeps SOZO A4 layout."""
    from docx.shared import Cm
    from sozo_generator.docx.renderer import DocumentRenderer
    from sozo_generator.template.doc_structure import build_document_spec
    from sozo_generator.core.enums import DocumentType, Tier

    spec = build_document_spec(
        condition=parkinsons_condition,
        doc_type=DocumentType.EVIDENCE_BASED_PROTOCOL,
        tier=Tier.FELLOW,
    )
    out_path = tmp_path / "default_layout.docx"
    renderer = DocumentRenderer(output_dir=str(tmp_path))
    renderer.render(spec, str(out_path))

    out_doc = Document(str(out_path))
    sec = out_doc.sections[0]
    assert sec.page_width.cm == pytest.approx(21.0, abs=_LEN_TOL_CM)
    assert sec.page_height.cm == pytest.approx(29.7, abs=_LEN_TOL_CM)


def test_template_driven_document_type_override(tmp_path: Path, parkinsons_condition):
    """Explicit document_type overrides filename inference."""
    from sozo_generator.template.template_driven_generator import TemplateDrivenGenerator
    from sozo_generator.core.enums import DocumentType, Tier

    layout_src = tmp_path / "evidence_protocol_template.docx"
    _write_non_a4_template(layout_src)

    gen = TemplateDrivenGenerator(layout_src)
    gen.parse_template()

    spec_inferred = gen.generate_for_condition(parkinsons_condition, Tier.FELLOW)
    assert spec_inferred.document_type == DocumentType.EVIDENCE_BASED_PROTOCOL

    spec_override = gen.generate_for_condition(
        parkinsons_condition,
        Tier.FELLOW,
        document_type=DocumentType.CLINICAL_EXAM,
    )
    assert spec_override.document_type == DocumentType.CLINICAL_EXAM


def _write_template_with_stationery_markers(path: Path) -> None:
    """Minimal template with recognizable header/footer text and one heading."""
    doc = Document()
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = "TPL_HDR_MARKER_STATIONERY"
    sec.footer.paragraphs[0].text = "TPL_FTR_MARKER_STATIONERY"
    doc.add_heading("Clinical Overview", level=1)
    doc.add_paragraph("Template body to clear.")
    doc.save(str(path))


def _header_and_footer_joined(doc: Document) -> tuple[str, str]:
    sec = doc.sections[0]
    header_text = " ".join(p.text for p in sec.header.paragraphs)
    footer_text = " ".join(p.text for p in sec.footer.paragraphs)
    return header_text, footer_text


def test_preserve_template_stationery_keeps_word_header_footer(
    parkinsons_condition, tmp_path: Path
):
    from sozo_generator.docx.renderer import DocumentRenderer
    from sozo_generator.template.doc_structure import build_document_spec
    from sozo_generator.core.enums import DocumentType, Tier

    tpl = tmp_path / "tpl_stationery.docx"
    _write_template_with_stationery_markers(tpl)

    spec = build_document_spec(
        condition=parkinsons_condition,
        doc_type=DocumentType.EVIDENCE_BASED_PROTOCOL,
        tier=Tier.FELLOW,
    )
    out_path = tmp_path / "out_stationery.docx"
    renderer = DocumentRenderer(output_dir=str(tmp_path))
    renderer.render(
        spec,
        str(out_path),
        layout_template_path=tpl,
        preserve_template_stationery=True,
    )

    out = Document(str(out_path))
    hdr, ftr = _header_and_footer_joined(out)
    assert "TPL_HDR_MARKER_STATIONERY" in hdr
    assert "TPL_FTR_MARKER_STATIONERY" in ftr
    assert "SOZO BRAIN CENTER" not in hdr


def test_layout_only_mode_replaces_template_header_with_soz(
    parkinsons_condition, tmp_path: Path
):
    from sozo_generator.docx.renderer import DocumentRenderer
    from sozo_generator.template.doc_structure import build_document_spec
    from sozo_generator.core.enums import DocumentType, Tier

    tpl = tmp_path / "tpl_replace.docx"
    _write_template_with_stationery_markers(tpl)

    spec = build_document_spec(
        condition=parkinsons_condition,
        doc_type=DocumentType.EVIDENCE_BASED_PROTOCOL,
        tier=Tier.FELLOW,
    )
    out_path = tmp_path / "out_replace.docx"
    renderer = DocumentRenderer(output_dir=str(tmp_path))
    renderer.render(
        spec,
        str(out_path),
        layout_template_path=tpl,
        preserve_template_stationery=False,
    )

    out = Document(str(out_path))
    hdr, _ = _header_and_footer_joined(out)
    assert "TPL_HDR_MARKER_STATIONERY" not in hdr
    assert "SOZO BRAIN CENTER" in hdr


def test_preserve_stationery_requires_layout_template(parkinsons_condition, tmp_path: Path):
    from sozo_generator.docx.renderer import DocumentRenderer
    from sozo_generator.template.doc_structure import build_document_spec
    from sozo_generator.core.enums import DocumentType, Tier

    spec = build_document_spec(
        condition=parkinsons_condition,
        doc_type=DocumentType.EVIDENCE_BASED_PROTOCOL,
        tier=Tier.FELLOW,
    )
    renderer = DocumentRenderer(output_dir=str(tmp_path))
    with pytest.raises(ValueError, match="preserve_template_stationery requires"):
        renderer.render(
            spec,
            str(tmp_path / "x.docx"),
            preserve_template_stationery=True,
        )
