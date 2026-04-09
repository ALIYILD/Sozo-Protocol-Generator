"""Authorization rules for protected /api routes (roles + authentication)."""
from __future__ import annotations

import asyncio
import uuid

import pytest

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient

from sozo_auth.tokens import create_access_token


def _bearer(role: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token('test-user', role)}"}

def _bearer_sub(sub: str, role: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token(sub, role)}"}

@pytest.fixture(autouse=True, scope="module")
def _ensure_graph_runs_table():
    """Ensure SQLAlchemy graph_runs table exists for graph authorization tests."""
    from sozo_db.base import Base
    from sozo_db.engine import get_engine

    async def _run() -> None:
        engine = get_engine()
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)

    asyncio.run(_run())


def _insert_graph_run(thread_id: str, created_by: uuid.UUID) -> None:
    from sozo_db.engine import get_session_factory
    from sozo_db.models.graph_run import GraphRun

    async def _run() -> None:
        factory = get_session_factory()
        async with factory() as session:
            session.add(
                GraphRun(
                    thread_id=thread_id,
                    status="queued",
                    condition_slug="test",
                    condition_name="Test",
                    source_mode="prompt",
                    created_by=created_by,
                    final_state={"request_id": thread_id, "status": "queued"},
                    node_history=[],
                    errors=[],
                )
            )
            await session.commit()

    asyncio.run(_run())

def _minimal_template_docx_bytes() -> bytes:
    """Tiny DOCX with one Heading for template-batch tests."""
    import io

    from docx import Document
    from docx.shared import Cm

    buf = io.BytesIO()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.59)
    sec.page_height = Cm(27.94)
    doc.add_heading("Clinical Overview", level=1)
    doc.add_paragraph("x")
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True, scope="module")
def _ensure_audit_log_table():
    """Create audit_log table in the test DB if it doesn't already exist."""
    from sozo_api.routes.db_helper import get_db
    conn = get_db()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS audit_log (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            entity_type TEXT    NOT NULL,
            entity_id   TEXT    NOT NULL,
            action      TEXT    NOT NULL,
            actor_id    TEXT,
            timestamp   TEXT    NOT NULL DEFAULT (datetime('now')),
            input_hash  TEXT,
            output_hash TEXT,
            node_name   TEXT,
            details     TEXT
        )
    """)
    conn.execute(
        "CREATE INDEX IF NOT EXISTS ix_audit_log_entity_ts "
        "ON audit_log (entity_type, entity_id, timestamp)"
    )
    conn.commit()
    conn.close()


@pytest.fixture
def client():
    from sozo_api.server import app

    return TestClient(app)


class TestProtocolsRouterAuth:
    def test_list_requires_authentication(self, client: TestClient):
        r = client.get("/api/protocols/")
        assert r.status_code in (401, 403)

    def test_list_allows_readonly(self, client: TestClient):
        r = client.get("/api/protocols/", headers=_bearer("readonly"))
        assert r.status_code == 200

    def test_create_requires_clinician(self, client: TestClient):
        body = {
            "condition_slug": "depression",
            "modality": "tdcs",
            "tier": "fellow",
            "doc_type": "evidence_based_protocol",
        }
        r = client.post("/api/protocols/", json=body, headers=_bearer("readonly"))
        assert r.status_code == 403

        r2 = client.post("/api/protocols/", json=body, headers=_bearer("clinician"))
        assert r2.status_code != 403

    def test_status_transition_requires_reviewer(self, client: TestClient):
        import uuid

        pid = str(uuid.uuid4())
        r = client.post(
            f"/api/protocols/{pid}/transition",
            json={"target_status": "approved", "notes": "sign-off"},
            headers=_bearer("clinician"),
        )
        assert r.status_code == 403


class TestPatientsRouterAuth:
    def test_requires_clinician_not_readonly(self, client: TestClient):
        r = client.get("/api/patients/", headers=_bearer("readonly"))
        assert r.status_code == 403

        r2 = client.get("/api/patients/", headers=_bearer("clinician"))
        assert r2.status_code != 403


class TestReviewsRouterAuth:
    def test_queue_requires_reviewer(self, client: TestClient):
        r = client.get("/api/reviews/queue", headers=_bearer("clinician"))
        assert r.status_code == 403

        r2 = client.get("/api/reviews/queue", headers=_bearer("reviewer"))
        assert r2.status_code == 200


class TestAuditRouterAuth:
    def test_events_requires_view_audit_permission(self, client: TestClient):
        r = client.get("/api/audit/events", headers=_bearer("clinician"))
        assert r.status_code == 403

        r2 = client.get("/api/audit/events", headers=_bearer("admin"))
        assert r2.status_code == 200

        r3 = client.get("/api/audit/events", headers=_bearer("operator"))
        assert r3.status_code == 200


class TestTemplateBatchAuth:
    def test_template_batch_requires_clinician(self, client: TestClient):
        docx = _minimal_template_docx_bytes()
        files = {
            "template": (
                "tpl.docx",
                docx,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        data = {"condition_slugs": "parkinsons", "tier": "fellow"}
        assert client.post("/api/generate/template-batch", files=files, data=data).status_code in (401, 403)
        r_ro = client.post(
            "/api/generate/template-batch",
            files=files,
            data=data,
            headers=_bearer("readonly"),
        )
        assert r_ro.status_code == 403
        r_ok = client.post(
            "/api/generate/template-batch",
            files=files,
            data=data,
            headers=_bearer("clinician"),
        )
        assert r_ok.status_code == 200
        assert r_ok.content[:2] == b"PK"


class TestInlineRoutesRoles:
    def test_cockpit_requires_operator_or_admin(self, client: TestClient):
        r = client.get("/api/cockpit/overview", headers=_bearer("clinician"))
        assert r.status_code == 403
        r2 = client.get("/api/cockpit/overview", headers=_bearer("operator"))
        assert r2.status_code == 200

    def test_evidence_staleness_requires_operator_or_admin(self, client: TestClient):
        r = client.get("/api/evidence/staleness", headers=_bearer("clinician"))
        assert r.status_code == 403
        r2 = client.get("/api/evidence/staleness", headers=_bearer("admin"))
        assert r2.status_code == 200

    def test_visuals_render_requires_clinician(self, client: TestClient):
        r = client.post("/api/visuals/render", json={}, headers=_bearer("readonly"))
        assert r.status_code == 403

    def test_graph_generate_requires_clinician(self, client: TestClient):
        r = client.post(
            "/api/graph/generate",
            json={"doc_type": "evidence_based_protocol", "tier": "fellow"},
            headers=_bearer("readonly"),
        )
        assert r.status_code == 403

    def test_graph_review_requires_reviewer(self, client: TestClient):
        r = client.post(
            "/api/graph/review",
            json={
                "thread_id": "00000000-0000-0000-0000-000000000000",
                "decision": "approve",
                "reviewer_id": "r1",
            },
            headers=_bearer("clinician"),
        )
        assert r.status_code == 403

    def test_graph_link_protocol_requires_clinician(self, client: TestClient):
        r = client.post(
            "/api/graph/link-protocol",
            json={
                "thread_id": "00000000-0000-0000-0000-000000000001",
                "protocol_id": "550e8400-e29b-41d4-a716-446655440000",
            },
            headers=_bearer("readonly"),
        )
        assert r.status_code == 403


class TestPublicReferenceRoutes:
    def test_health_and_knowledge_unauthenticated(self, client: TestClient):
        assert client.get("/api/health").status_code == 200
        assert client.get("/api/knowledge/conditions").status_code == 200
        assert client.get("/api/visuals/types").status_code == 200


class TestRefreshTokenType:
    def test_refresh_endpoint_rejects_access_token(self, client: TestClient):
        from sozo_auth.tokens import create_access_token

        access = create_access_token("user-refresh-test", "clinician")
        r = client.post("/api/auth/refresh", json={"refresh_token": access})
        assert r.status_code == 401
        assert r.json().get("detail") == "Refresh token required"

    def test_bearer_refresh_token_rejected_for_api(self, client: TestClient):
        from sozo_auth.tokens import create_refresh_token

        refresh = create_refresh_token("user-bearer-test")
        r = client.get(
            "/api/protocols/",
            headers={"Authorization": f"Bearer {refresh}"},
        )
        assert r.status_code == 401


class TestGraphThreadOwnership:
    def test_graph_status_forbidden_for_non_owner(self, client: TestClient):
        thread_id = str(uuid.uuid4())
        owner_id = uuid.uuid4()
        other_user_id = uuid.uuid4()
        _insert_graph_run(thread_id, created_by=owner_id)

        r = client.get(
            f"/api/graph/status/{thread_id}",
            headers=_bearer_sub(str(other_user_id), "clinician"),
        )
        assert r.status_code == 403

    def test_graph_link_protocol_forbidden_for_non_owner(self, client: TestClient):
        thread_id = str(uuid.uuid4())
        owner_id = uuid.uuid4()
        other_user_id = uuid.uuid4()
        _insert_graph_run(thread_id, created_by=owner_id)

        r = client.post(
            "/api/graph/link-protocol",
            json={
                "thread_id": thread_id,
                "protocol_id": "550e8400-e29b-41d4-a716-446655440000",
            },
            headers=_bearer_sub(str(other_user_id), "clinician"),
        )
        assert r.status_code == 403

    def test_graph_review_forbidden_for_non_owner(self, client: TestClient):
        thread_id = str(uuid.uuid4())
        owner_id = uuid.uuid4()
        other_user_id = uuid.uuid4()
        _insert_graph_run(thread_id, created_by=owner_id)

        r = client.post(
            "/api/graph/review",
            json={
                "thread_id": thread_id,
                "decision": "approve",
                "reviewer_id": "ignored",
            },
            headers=_bearer_sub(str(other_user_id), "reviewer"),
        )
        assert r.status_code == 403
